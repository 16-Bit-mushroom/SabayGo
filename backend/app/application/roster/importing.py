"""Member-roster import (brief 8): operators and vans from the binder.

The office downloads a template, copies the binder into it and uploads
it. Every row is checked before anything is saved -- required fields,
plate / licence / CPC formats, no duplicate of a record on file or of
another row, one van per operator -- and a preview shows each row with
its reasons. Only a confirmed import writes. It re-reads the file and
checks it again rather than trusting a preview the roster may have moved
on from, saves the passing rows as Pending under one import batch, and
keeps the failing rows with their reasons as a CSV to correct and upload
again. Every import is logged in data_transfers.

Dispatchers are not imported: each is a staff login, made one at a time.
The vans' four photographs are not either; they are uploaded afterwards.
"""

from __future__ import annotations

import csv
import datetime as dt
import io
import uuid
import zipfile
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.db import commit_or_conflict
from app.core.exceptions import InvalidInputError
from app.core.timezone import now
from app.domain.value_objects import (
    CpcReference,
    Email,
    LicenseNumber,
    PhoneNumber,
    PlateNumber,
)
from app.infrastructure.models import DataTransfer, Operator, Route, User, Van

IMPORTABLE = ("operators", "vans")

# Template columns, in order. A van names its operator by licence number:
# the one identifier on both the operator's binder page and the van's.
COLUMNS: dict[str, tuple[str, ...]] = {
    "operators": ("full_name", "contact_number", "email", "license_number",
                  "license_expiry"),
    "vans": ("plate_number", "operator_license_number", "cpc_case_no",
             "cpc_number", "brand", "model", "color", "seat_capacity",
             "registered_route"),
}
REQUIRED: dict[str, frozenset[str]] = {
    "operators": frozenset({"full_name", "contact_number", "license_number",
                            "license_expiry"}),
    "vans": frozenset({"plate_number", "operator_license_number"}),
}
HELP: dict[str, str] = {
    "full_name": "As on the licence",
    "contact_number": "Mobile number, 09XXXXXXXXX",
    "email": "Optional",
    "license_number": "LTO licence, e.g. N01-23-456789",
    "license_expiry": "YYYY-MM-DD (or MM/DD/YYYY)",
    "plate_number": "e.g. ABC-1234",
    "operator_license_number": "Licence number of the van's operator, already on the roster",
    "cpc_case_no": "LTFRB CPC case number",
    "cpc_number": "LTFRB CPC number",
    "brand": "e.g. Toyota",
    "model": "e.g. HiAce Commuter",
    "color": "e.g. White",
    "seat_capacity": "1 to 14; blank means 14",
    "registered_route": "Route code, e.g. ECO-TAG",
}
_TEXT_COLUMNS = {"contact_number", "license_number", "plate_number",
                 "operator_license_number", "cpc_case_no", "cpc_number"}
_LIMITS = {"full_name": 255, "email": 255, "brand": 64, "model": 64, "color": 32}

MAX_BYTES = 2 * 1024 * 1024
MAX_ROWS = 2000
REJECTED_DIR = Path(__file__).resolve().parents[3] / "storage" / "transfers"


def rejected_path(transfer_id: str) -> Path:
    return REJECTED_DIR / f"{transfer_id}-rejected.csv"


# ------------------------------------------------------------------ files
def template(record_type: str, fmt: str) -> bytes:
    columns = COLUMNS[record_type]
    if fmt == "csv":
        buf = io.StringIO()
        csv.writer(buf).writerow(columns)
        return buf.getvalue().encode("utf-8-sig")

    wb = Workbook()
    ws = wb.active
    ws.title = record_type
    ws.append(list(columns))
    for i, col in enumerate(columns, start=1):
        letter = ws.cell(row=1, column=i).column_letter
        ws.column_dimensions[letter].width = max(16, len(col) + 2)
        if col in _TEXT_COLUMNS:
            # Text, or Excel drops a mobile number's leading zero.
            for row in range(2, 1002):
                ws.cell(row=row, column=i).number_format = "@"
    notes = wb.create_sheet("How to fill")
    notes.append(["column", "required", "what to enter"])
    for col in columns:
        notes.append([col, "yes" if col in REQUIRED[record_type] else "no", HELP[col]])
    notes.column_dimensions["A"].width = 26
    notes.column_dimensions["C"].width = 60
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _cell(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, dt.datetime):
        return value.date().isoformat()
    if isinstance(value, dt.date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _read_table(filename: str, data: bytes) -> tuple[str, list[list[str]]]:
    if len(data) > MAX_BYTES:
        raise InvalidInputError("The file is larger than 2 MB.")
    name = filename.lower()
    if name.endswith(".csv"):
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            raise InvalidInputError(
                "Save the file as 'CSV UTF-8' and upload it again."
            ) from None
        return "csv", [[c.strip() for c in row] for row in csv.reader(io.StringIO(text))]
    if name.endswith(".xlsx"):
        try:
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except (zipfile.BadZipFile, InvalidFileException, KeyError, ValueError, OSError):
            raise InvalidInputError("This is not a readable Excel (.xlsx) file.") from None
        table: list[list[str]] = []
        # A sheet can declare a million formatted rows; stop well short.
        for i, row in enumerate(wb.worksheets[0].iter_rows(values_only=True)):
            if i > MAX_ROWS * 5:
                break
            table.append([_cell(c) for c in row])
        wb.close()
        return "xlsx", table
    raise InvalidInputError("Upload the CSV or Excel (.xlsx) template.")


def read_rows(record_type: str, filename: str, data: bytes) -> tuple[str, list[tuple[int, dict[str, str]]]]:
    """(file format, [(spreadsheet row number, values by column)]).

    Columns are matched by name, ignoring case and spaces; extra columns
    -- the 'reason' of a rejected-rows file -- are ignored, so that file
    can be corrected and uploaded as it is. Blank rows are skipped.
    """
    fmt, table = _read_table(filename, data)
    if not table:
        raise InvalidInputError("The file is empty.")
    header = [h.strip().lower().replace(" ", "_") for h in table[0]]
    missing = [c for c in COLUMNS[record_type]
               if c in REQUIRED[record_type] and c not in header]
    if missing:
        raise InvalidInputError(
            f"The file has no column {', '.join(missing)}. Use the {record_type} template."
        )
    index = {c: header.index(c) for c in COLUMNS[record_type] if c in header}
    rows = []
    for number, raw in enumerate(table[1:], start=2):
        if not any(raw):
            continue
        rows.append((number, {c: raw[i] if i < len(raw) else "" for c, i in index.items()}))
    if len(rows) > MAX_ROWS:
        raise InvalidInputError(f"The file has more than {MAX_ROWS} rows.")
    return fmt, rows


# ----------------------------------------------------------------- checks
@dataclass
class RowResult:
    row: int
    values: dict[str, str]
    errors: list[str] = field(default_factory=list)
    record: dict | None = None  # the cleaned fields, when the row passes

    def as_dict(self) -> dict:
        return {"row": self.row, "ok": not self.errors,
                "errors": self.errors, "values": self.values}


def _date(raw: str) -> dt.date:
    for pattern in ("%Y-%m-%d", "%m/%d/%Y"):
        try:
            return dt.datetime.strptime(raw, pattern).date()
        except ValueError:
            pass
    raise ValueError("must be a date, YYYY-MM-DD")


def _phone(raw: str) -> str:
    # A number cell in Excel loses the leading zero: 9171234567.
    if raw.isdigit() and len(raw) == 10 and raw.startswith("9"):
        raw = "0" + raw
    return PhoneNumber(raw).normalized()


class _Row:
    """One row's checks; collects every reason rather than the first."""

    def __init__(self, number: int, values: dict[str, str], required: frozenset[str]):
        self.number = number
        self.values = values
        self.required = required
        self.errors: list[str] = []

    def get(self, col: str, parse=None):
        raw = self.values.get(col, "").strip()
        if not raw:
            if col in self.required:
                self.errors.append(f"{col} is required")
            return None
        if col in _LIMITS and len(raw) > _LIMITS[col]:
            self.errors.append(f"{col} is longer than {_LIMITS[col]} characters")
            return None
        if parse is None:
            return raw
        try:
            return parse(raw)
        except ValueError as exc:
            self.errors.append(f"{col}: {exc}")
            return None

    def unique(self, col: str, value, on_file: set, seen: dict) -> None:
        """Not already on the roster, and not on another row of this file."""
        if value is None:
            return
        if value in on_file:
            self.errors.append(f"{col} {value} is already on the roster")
        elif value in seen:
            self.errors.append(f"{col} {value} is also on row {seen[value]}")
        seen.setdefault(value, self.number)


async def check_rows(
    session: AsyncSession, record_type: str, rows: list[tuple[int, dict[str, str]]]
) -> list[RowResult]:
    """Check every row against the roster and against the rows above it.

    A key repeated within the file fails on the later row whatever else
    is wrong with the earlier one: fixing the earlier row would make
    them clash anyway.
    """
    results = []
    if record_type == "operators":
        contacts = set((await session.execute(select(Operator.contact_number))).scalars())
        licences = set((await session.execute(select(Operator.license_number))).scalars())
        seen_contact: dict = {}
        seen_licence: dict = {}
        for number, values in rows:
            r = _Row(number, values, REQUIRED[record_type])
            name = r.get("full_name")
            contact = r.get("contact_number", _phone)
            email = r.get("email", lambda v: Email(v).value)
            licence = r.get("license_number", lambda v: LicenseNumber(v).value)
            expiry = r.get("license_expiry", _date)
            r.unique("contact_number", contact, contacts, seen_contact)
            r.unique("license_number", licence, licences, seen_licence)
            record = None if r.errors else {
                "full_name": name, "contact_number": contact, "email": email,
                "license_number": licence, "license_expiry": expiry,
            }
            results.append(RowResult(number, values, r.errors, record))
        return results

    plates = set((await session.execute(select(Van.plate_number))).scalars())
    operators = {
        lic: (op_id, name, plate)
        for lic, op_id, name, plate in (await session.execute(
            select(Operator.license_number, Operator.operator_id,
                   Operator.full_name, Van.plate_number)
            .outerjoin(Van, Van.operator_id == Operator.operator_id)
        )).all()
    }
    routes = dict((await session.execute(select(Route.route_code, Route.route_id))).all())
    seen_plate: dict = {}
    seen_operator: dict = {}
    for number, values in rows:
        r = _Row(number, values, REQUIRED[record_type])
        plate = r.get("plate_number", lambda v: PlateNumber(v).value)
        licence = r.get("operator_license_number", lambda v: LicenseNumber(v).value)
        cpc_case = r.get("cpc_case_no", lambda v: CpcReference(v).value)
        cpc_no = r.get("cpc_number", lambda v: CpcReference(v).value)
        brand, model, color = r.get("brand"), r.get("model"), r.get("color")
        seats = r.get("seat_capacity", _seats)
        route_code = r.get("registered_route")
        r.unique("plate_number", plate, plates, seen_plate)

        operator_id = None
        if licence is not None:
            found = operators.get(licence)
            if found is None:
                r.errors.append(
                    f"no operator with licence {licence} is on the roster; import operators first"
                )
            elif found[2] is not None:
                r.errors.append(f"{found[1]} already has van {found[2]}: one van per operator")
            elif licence in seen_operator:
                r.errors.append(
                    f"{found[1]} is also named on row {seen_operator[licence]}: one van per operator"
                )
            else:
                operator_id = found[0]
            seen_operator.setdefault(licence, number)

        route_id = None
        if route_code is not None:
            route_id = routes.get(route_code.upper())
            if route_id is None:
                r.errors.append(f"no route with code {route_code}")

        record = None if r.errors else {
            "plate_number": plate, "operator_id": operator_id,
            "cpc_case_no": cpc_case, "cpc_number": cpc_no, "brand": brand,
            "model": model, "color": color, "seat_capacity": seats or 14,
            "registered_route_id": route_id,
        }
        results.append(RowResult(number, values, r.errors, record))
    return results


def _seats(raw: str) -> int:
    if not raw.isdigit() or not 1 <= int(raw) <= 14:
        raise ValueError("must be a whole number from 1 to 14")
    return int(raw)


# ----------------------------------------------------------------- import
async def preview(
    session: AsyncSession, record_type: str, filename: str, data: bytes
) -> list[RowResult]:
    _, rows = read_rows(record_type, filename, data)
    return await check_rows(session, record_type, rows)


async def commit(
    session: AsyncSession, record_type: str, filename: str, data: bytes, user: User
) -> tuple[DataTransfer, list[RowResult]]:
    """Save the passing rows as one Pending batch and log the import."""
    fmt, rows = read_rows(record_type, filename, data)
    results = await check_rows(session, record_type, rows)
    passing = [r for r in results if r.record is not None]
    failing = [r for r in results if r.record is None]

    stamp = now().replace(tzinfo=None)
    transfer = DataTransfer(
        transfer_id=str(uuid.uuid4()),
        direction="import",
        record_type=record_type,
        file_name=filename[:255],
        file_format=fmt,
        rows_in_file=len(results),
        rows_accepted=len(passing),
        rows_rejected=len(failing),
        performed_by=user.user_id,
        performed_at=stamp,
    )
    if failing:
        transfer.error_file_url = f"/api/v1/roster/transfers/{transfer.transfer_id}/rejected"
    session.add(transfer)
    await session.flush()  # the batch row first: records point at it

    model = Operator if record_type == "operators" else Van
    id_field = "operator_id" if record_type == "operators" else "van_id"
    extra = ({"membership_status": "active"} if record_type == "operators"
             else {"operational_status": "active"})
    for r in passing:
        session.add(model(
            **r.record, **extra,
            **{id_field: str(uuid.uuid4())},
            validation_status="pending",
            import_batch_id=transfer.transfer_id,
            created_at=stamp,
            updated_at=stamp,
        ))

    path = rejected_path(transfer.transfer_id)
    if failing:
        REJECTED_DIR.mkdir(parents=True, exist_ok=True)
        path.write_bytes(_rejected_csv(record_type, failing))
    try:
        await commit_or_conflict(
            session,
            "The roster changed while this file was being imported. Preview it again.",
        )
    except Exception:
        path.unlink(missing_ok=True)
        raise
    return transfer, results


def _rejected_csv(record_type: str, failing: list[RowResult]) -> bytes:
    columns = COLUMNS[record_type]
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow([*columns, "reason"])
    for r in failing:
        w.writerow([*(r.values.get(c, "") for c in columns), "; ".join(r.errors)])
    return buf.getvalue().encode("utf-8-sig")
